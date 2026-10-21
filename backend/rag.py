import os
import json
import base64
import csv
import io
import requests
from bs4 import BeautifulSoup
from docx import Document
from pptx import Presentation
import pandas as pd
from youtube_transcript_api import YouTubeTranscriptApi
import chromadb
from chromadb.utils import embedding_functions
from openai import AsyncOpenAI
from pypdf import PdfReader
import re
import asyncio
import edge_tts
import uuid
from pydub import AudioSegment
import urllib.parse
import aiohttp
from math_solver import solve_math_problem_sympy, solve_math_problem
from PIL import Image

# Initialize ChromaDB
chroma_client = chromadb.PersistentClient(path="./chroma_data")
try:
    from chromadb.utils.embedding_functions import SentenceTransformerEmbeddingFunction
    emb_fn = SentenceTransformerEmbeddingFunction(model_name="all-MiniLM-L6-v2")
except Exception:
    emb_fn = embedding_functions.DefaultEmbeddingFunction()

collection = chroma_client.get_or_create_collection(name="studysnap_collection", embedding_function=emb_fn)

# Initialize LM Studio client
LM_STUDIO_URL = "http://localhost:1234/v1"
lm_studio_client = AsyncOpenAI(base_url=LM_STUDIO_URL, api_key="lm-studio")

def get_available_models_sync() -> list:
    """Fetch active models from LM Studio endpoints."""
    try:
        resp = requests.get(f"{LM_STUDIO_URL}/models", timeout=2)
        if resp.status_code == 200:
            data = resp.json()
            return [m.get("id") for m in data.get("data", []) if m.get("id")]
    except Exception:
        pass
    return []

def get_best_model_name(is_vision: bool = False) -> str:
    """Dynamically determine the best model identifier from LM Studio."""
    models = get_available_models_sync()
    if is_vision:
        vision_models = [m for m in models if any(k in m.lower() for k in ["vl", "vision", "qwen3-vl", "llava"])]
        if vision_models:
            return vision_models[0]
    text_models = [m for m in models if "embed" not in m.lower()]
    if text_models:
        return text_models[0]
    if models:
        return models[0]
    return "qwen/qwen3-vl-8b"

def extract_youtube_transcript(url: str) -> str:
    try:
        video_id = ""
        if "v=" in url:
            video_id = url.split("v=")[1].split("&")[0]
        elif "youtu.be/" in url:
            video_id = url.split("youtu.be/")[1].split("?")[0]
        elif "embed/" in url:
            video_id = url.split("embed/")[1].split("?")[0]
            
        if not video_id: 
            raise ValueError("Could not find video ID in URL.")
        
        ytt_api = YouTubeTranscriptApi()
        
        try:
            transcript_list = ytt_api.list(video_id)
            transcript = transcript_list.find_transcript(['en', 'en-US', 'en-GB', 'hi'])
            data = transcript.fetch()
        except:
            try:
                transcript_list = ytt_api.list(video_id)
                transcript = list(transcript_list)[0]
                data = transcript.fetch()
            except Exception as e:
                data = ytt_api.fetch(video_id)
                
        return " ".join([t.text if hasattr(t, 'text') else t.get('text', '') for t in data])
    except Exception as e:
        print(f"YouTube Extraction Error: {e}")
        raise ValueError(f"Failed to process YouTube transcript: {str(e)}. The video might not have captions enabled.")

def extract_webpage(url: str) -> str:
    try:
        headers = {"User-Agent": "Mozilla/5.0"}
        response = requests.get(url, headers=headers)
        soup = BeautifulSoup(response.text, 'html.parser')
        for script in soup(["script", "style"]):
            script.extract()
        return soup.get_text(separator="\n", strip=True)
    except Exception as e:
        print(f"Webpage Extraction Error: {e}")
        return ""

try:
    import whisper
except Exception as _w_err:
    whisper = None
    print(f"Whisper import warning: {_w_err}")

_easyocr_reader = None

def get_ocr_reader():
    global _easyocr_reader
    if _easyocr_reader is None:
        try:
            import easyocr
            _easyocr_reader = easyocr.Reader(['en'], gpu=False)
        except Exception as e:
            print(f"EasyOCR Init Warning: {e}")
            _easyocr_reader = False
    return _easyocr_reader

def format_image_analysis_to_markdown(analysis: dict) -> str:
    filename = analysis.get("filename", "image")
    md_parts = [f"IMAGE SOURCE: {filename}"]
    
    if desc := analysis.get("image_description"):
        md_parts.append(f"### IMAGE OVERVIEW & DESCRIPTION:\n{desc}")
        
    if elem := analysis.get("visual_elements"):
        val_str = "\n".join([f"- {item}" for item in elem]) if isinstance(elem, list) else str(elem)
        md_parts.append(f"### VISUAL ELEMENTS & DIAGRAM COMPONENTS:\n{val_str}")
        
    if rel := analysis.get("relationships"):
        val_str = "\n".join([f"- {item}" for item in rel]) if isinstance(rel, list) else str(rel)
        md_parts.append(f"### RELATIONSHIPS & DIAGRAM FLOW:\n{val_str}")

    if text := analysis.get("extracted_text"):
        md_parts.append(f"### EXTRACTED TEXT & LABELS:\n{text}")
        
    if info := analysis.get("important_information"):
        val_str = "\n".join([f"- {item}" for item in info]) if isinstance(info, list) else str(info)
        md_parts.append(f"### IMPORTANT CONCEPTS & SUMMARY:\n{val_str}")
        
    if math_sol := analysis.get("math_solution"):
        md_parts.append(f"### VERIFIED MATHEMATICAL SOLUTION:\n{math_sol}")
        
    return "\n\n".join(md_parts)

async def extract_image_content_async(file_path: str, filename: str, file_id: str = "") -> str:
    """
    Extract visual description, diagram components, relationships, OCR text,
    and math from an image using Vision LLM + OCR + SymPy.
    Caches structured JSON results to disk for fast reuse.
    """
    if not file_path or not os.path.exists(file_path):
        return f"IMAGE SOURCE: {filename}\n### IMAGE OVERVIEW & DESCRIPTION:\nImage file provided for study context."

    cache_path = f"{file_path}.analysis.json"
    if os.path.exists(cache_path):
        try:
            with open(cache_path, "r", encoding="utf-8") as f:
                analysis = json.load(f)
                return format_image_analysis_to_markdown(analysis)
        except Exception as e:
            print(f"Error reading image analysis cache for {filename}: {e}")

    extracted_text = ""
    image_description = ""
    visual_elements = ""
    relationships = ""
    important_info = ""
    math_solution = ""
    detected_math = ""

    # 1. Vision LLM Analysis via LM Studio (Fast & Full Visual Understanding)
    try:
        with open(file_path, "rb") as f_img:
            b64_str = base64.b64encode(f_img.read()).decode("utf-8")

        ext_mime = filename.split('.')[-1].lower() if '.' in filename else "png"
        mime_type = "image/png" if ext_mime in ['png', 'webp'] else "image/jpeg"
        vision_model = get_best_model_name(is_vision=True)

        vision_prompt = (
            "Analyze this image thoroughly for an educational RAG study assistant.\n"
            "Return a JSON object with EXACTLY the following keys:\n"
            "{\n"
            '  "image_description": "Comprehensive explanation of what the image, diagram, architecture, or flowchart represents.",\n'
            '  "extracted_text": "All visible titles, headings, node labels, text, and numbers in the image.",\n'
            '  "visual_elements": "List of components, boxes, icons, servers, routers, clients, tables, or visual nodes.",\n'
            '  "relationships": "Step-by-step description of connections, arrows, sequence, or workflow shown in the diagram.",\n'
            '  "important_information": "Key educational concepts, summary points, or architectural takeaways.",\n'
            '  "detected_math": "Any mathematical formulas or equations present in the image (or empty string if none)."\n'
            "}\n"
            "Return strictly valid JSON without extra text."
        )

        response = await lm_studio_client.chat.completions.create(
            model=vision_model,
            messages=[
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": vision_prompt},
                        {
                            "type": "image_url",
                            "image_url": {
                                "url": f"data:{mime_type};base64,{b64_str}"
                            }
                        }
                    ]
                }
            ],
            temperature=0.2,
            max_tokens=1500
        )

        raw_content = response.choices[0].message.content.strip()
        parsed = {}
        try:
            json_match = re.search(r'\{.*\}', raw_content, re.DOTALL)
            if json_match:
                parsed = json.loads(json_match.group(0))
            else:
                parsed = json.loads(raw_content)
        except Exception:
            image_description = raw_content

        if parsed:
            image_description = parsed.get("image_description", image_description)
            extracted_text = parsed.get("extracted_text", "")
            visual_elements = parsed.get("visual_elements", "")
            relationships = parsed.get("relationships", "")
            important_info = parsed.get("important_information", "")
            detected_math = parsed.get("detected_math", "")

    except Exception as v_err:
        print(f"Vision LLM analysis notice for {filename}: {v_err}")

    # 2. OCR Fallback if Vision LLM was offline or text missing
    if not extracted_text:
        try:
            with Image.open(file_path) as img:
                try:
                    import pytesseract
                    ocr_result = pytesseract.image_to_string(img).strip()
                    if ocr_result:
                        extracted_text = ocr_result
                except Exception:
                    pass

                if not extracted_text:
                    try:
                        reader = get_ocr_reader()
                        if reader and reader is not False:
                            results = reader.readtext(file_path, detail=0)
                            if results:
                                extracted_text = "\n".join(results).strip()
                    except Exception as e:
                        print(f"EasyOCR extraction error: {e}")
        except Exception as e:
            print(f"Image load error: {e}")

    # 3. Fallback heuristic synthesis if Vision LLM was offline
    if not image_description:
        lines = [line.strip() for line in extracted_text.splitlines() if line.strip()]
        if lines:
            image_description = f"Visual diagram / image source '{filename}' titled '{lines[0]}'."
            visual_elements = "Extracted components & text labels: " + ", ".join(lines)
            flow_steps = [f"{idx}. {line}" for idx, line in enumerate(lines, 1)]
            relationships = "Diagram Flow & Components:\n" + "\n".join(flow_steps)
            important_info = f"Key concepts from '{filename}': " + ", ".join(lines[:6])
        else:
            image_description = f"Visual diagram / illustration source '{filename}'."
            visual_elements = f"Graphical layout and components present in image '{filename}'."
            relationships = f"Visual layout in '{filename}'."
            important_info = f"Visual study diagram from '{filename}'."

    # 4. Math Solver Integration if math detected
    combined_math_target = f"{extracted_text}\n{detected_math}"
    if combined_math_target and any(c in combined_math_target for c in ['=', '+', '-', '*', '/', '^', '√', 'd/dx', '∫']):
        try:
            math_res = solve_math_problem_sympy(combined_math_target)
            if math_res.get("steps") and math_res.get("verification", {}).get("status") == "verified":
                math_solution = format_sympy_result_to_markdown(math_res)
        except Exception as m_err:
            print(f"Math solver error on image content: {m_err}")

    # 5. Save Structured Representation to Cache
    analysis_data = {
        "source_type": "image",
        "filename": filename,
        "file_id": file_id,
        "extracted_text": extracted_text,
        "image_description": image_description,
        "visual_elements": visual_elements,
        "relationships": relationships,
        "important_information": important_info,
        "math_solution": math_solution
    }

    try:
        with open(cache_path, "w", encoding="utf-8") as f_cache:
            json.dump(analysis_data, f_cache, indent=2)
    except Exception as c_err:
        print(f"Error saving image analysis cache: {c_err}")

    return format_image_analysis_to_markdown(analysis_data)

def extract_image_content(file_path: str, filename: str, file_id: str = "") -> str:
    """Synchronous wrapper for extract_image_content_async."""
    try:
        loop = asyncio.get_event_loop()
        if loop.is_running():
            import nest_asyncio
            nest_asyncio.apply()
            return loop.run_until_complete(extract_image_content_async(file_path, filename, file_id))
        else:
            return asyncio.run(extract_image_content_async(file_path, filename, file_id))
    except Exception:
        return asyncio.run(extract_image_content_async(file_path, filename, file_id))

def extract_audio(file_path: str) -> str:
    """Extract text from audio using local Whisper instance."""
    try:
        if whisper is None:
            return "Whisper library is not available."
        model = whisper.load_model("base")
        result = model.transcribe(file_path)
        return result.get("text", "")
    except Exception as e:
        print(f"Whisper Audio Extraction Error: {e}")
        return "Audio transcription failed. Please ensure ffmpeg is installed."

def extract_text(file_path: str, filename: str) -> str:
    """Extract text from all supported file types."""
    text = ""
    ext = filename.split('.')[-1].lower() if '.' in filename else ""
    
    try:
        if ext in ['jpg', 'jpeg', 'png', 'webp']:
            text = extract_image_content(file_path, filename)
        elif ext == 'pdf':
            reader = PdfReader(file_path)
            for page in reader.pages:
                if page_text := page.extract_text(): text += page_text + "\n"
        elif ext in ['txt', 'md', 'html', 'xml']:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
                if ext in ['html', 'xml']:
                    soup = BeautifulSoup(content, 'html.parser')
                    text = soup.get_text(separator="\n", strip=True)
                else:
                    text = content
        elif ext == 'json':
            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
                text = json.dumps(data, indent=2)
        elif ext == 'csv':
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                reader = csv.reader(f)
                for row in reader:
                    text += " ".join(row) + "\n"
        elif ext == 'xlsx':
            df = pd.read_excel(file_path)
            text = df.to_string()
        elif ext == 'docx':
            doc = Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs])
        elif ext == 'pptx':
            prs = Presentation(file_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
        elif ext in ['wav', 'mp3', 'm4a', 'flac']:
            text = extract_audio(file_path)
    except Exception as e:
        print(f"Error extracting {filename}: {e}")
        
    return text

def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 200):
    chunks = []
    start = 0
    text_length = len(text)
    while start < text_length:
        end = start + chunk_size
        chunks.append(text[start:end])
        start = end - overlap
    return chunks

async def ingest_document(file_path: str, filename: str, file_id: str):
    """Ingest document/image/URL into ChromaDB vector store."""
    text = ""
    ext = filename.split('.')[-1].lower() if '.' in filename else ""
    is_image = ext in ['jpg', 'jpeg', 'png', 'webp'] or (file_path and file_path.lower().endswith(('jpg', 'jpeg', 'png', 'webp')))
    
    if filename.startswith("http"):
        if "youtube.com" in filename or "youtu.be" in filename:
            text = extract_youtube_transcript(filename)
            filename = f"YouTube: {filename}"
        else:
            text = extract_webpage(filename)
            filename = f"Web: {filename}"
    elif is_image:
        text = await extract_image_content_async(file_path, filename, file_id)
    else:
        text = extract_text(file_path, filename)
        
    if not text.strip():
        raise ValueError("Could not extract any content from the uploaded file.")

    chunks = chunk_text(text)
    ids = [f"{file_id}_chunk_{i}" for i in range(len(chunks))]
    
    source_type = "image" if is_image else ("url" if filename.startswith("http") else (ext if ext else "file"))
    file_ref = os.path.basename(file_path) if file_path else ""
    image_url = f"http://localhost:8000/uploads/{file_ref}" if is_image and file_ref else ""
    
    metadatas = [{
        "source": filename,
        "file_id": file_id,
        "source_type": source_type,
        "type": source_type,
        "filename": filename,
        "url": image_url
    } for _ in chunks]
    
    collection.add(documents=chunks, metadatas=metadatas, ids=ids)
    return {
        "id": file_id,
        "filename": filename,
        "type": source_type,
        "source_type": source_type,
        "url": image_url
    }

def get_all_sources():
    results = collection.get(include=["metadatas"])
    metadatas = results.get("metadatas", [])
    unique_sources = {}
    for meta in metadatas:
        fid = meta.get("file_id")
        if fid and fid not in unique_sources:
            filename = meta.get("source", "")
            ext = filename.split('.')[-1].lower() if '.' in filename and not filename.startswith("http") else "file"
            stype = meta.get("source_type") or meta.get("type") or ("image" if ext in ['jpg', 'jpeg', 'png', 'webp'] else ext)
            url = meta.get("url", "")
            
            if stype == "image" and not url:
                try:
                    for fname in os.listdir("uploads"):
                        if fname.startswith(fid):
                            url = f"http://localhost:8000/uploads/{fname}"
                            break
                except Exception:
                    pass
                    
            unique_sources[fid] = {
                "id": fid,
                "filename": filename,
                "type": stype,
                "source_type": stype,
                "url": url
            }
    return list(unique_sources.values())

def delete_source(file_id: str):
    """Delete a source and all its chunks from ChromaDB."""
    collection.delete(where={"file_id": file_id})

def format_sympy_result_to_markdown(res: dict) -> str:
    md = f"### 📐 {res.get('problem_type', 'Math Problem Solver')}\n\n"
    if res.get('concept'):
        md += f"**Concept**: {res.get('concept')}\n\n"
    
    if res.get("given"):
        md += "**Given:**\n"
        for g in res["given"]:
            md += f"- {g}\n"
        md += "\n"
        
    if res.get("required"):
        md += f"**Required:** {res['required']}\n\n"
        
    if res.get("formulas"):
        md += "**Formulas:**\n"
        for f in res["formulas"]:
            latex_val = f.get('latex', '')
            if latex_val:
                if not (latex_val.startswith('$') or latex_val.startswith('\\(')):
                    latex_val = f"$${latex_val}$$"
                md += f"- **{f['name']}**:\n{latex_val}\n"
            else:
                md += f"- **{f['name']}**\n"
        md += "\n"
        
    if res.get("steps"):
        md += "**Step-by-Step Solution:**\n\n"
        for step in res["steps"]:
            md += f"#### Step {step.get('step', 1)}: {step.get('explanation', '')}\n"
            if step.get("formula"):
                f_str = str(step['formula'])
                if not (f_str.startswith('$') or f_str.startswith('\\(')):
                    f_str = f"${f_str}$"
                md += f"- **Formula**: {f_str}\n"
            if step.get("calculation"):
                c_str = str(step['calculation'])
                if not (c_str.startswith('$') or c_str.startswith('\\(')):
                    c_str = f"${c_str}$"
                md += f"- **Calculation**: {c_str}\n"
            if step.get("result"):
                r_str = str(step['result'])
                if not (r_str.startswith('$') or r_str.startswith('\\(')):
                    r_str = f"${r_str}$"
                md += f"- **Result**: {r_str}\n\n"
        
    if res.get("final_answer"):
        ans_str = str(res['final_answer'])
        if not (ans_str.startswith('$') or ans_str.startswith('\\(')):
            ans_str = f"$${ans_str}$$"
        md += f"### ✅ **Final Answer:**\n\n{ans_str}\n\n"
        
    if res.get("verification", {}).get("status") == "verified":
        md += f"🔍 *Verification*: Verified mathematically using SymPy Engine ({res['verification'].get('result', '')}).\n"
    elif res.get("explanation"):
        md += f"\n*{res['explanation']}*\n"
        
    return md

async def chat_with_context(query: str, history: list = None, selected_source_ids: list = None, response_language: str = "English"):
    if history is None: history = []
    
    # 1. First check if query is a deterministic math problem
    sympy_res = solve_math_problem_sympy(query)
    is_verified_math = sympy_res.get("steps") and sympy_res.get("verification", {}).get("status") == "verified"
    
    if is_verified_math:
        math_md = format_sympy_result_to_markdown(sympy_res)
        try:
            active_model = get_best_model_name(is_vision=False)
            enriched_res = await solve_math_problem(query, lm_studio_client, active_model)
            if isinstance(enriched_res, dict) and enriched_res.get("steps"):
                return format_sympy_result_to_markdown(enriched_res), ["SymPy Engine & Local LLM"]
        except Exception:
            pass
        return math_md, ["SymPy Math Engine"]
        
    # 2. Document & Image RAG Retrieval
    where_clause = None
    if selected_source_ids is not None and len(selected_source_ids) > 0:
        if len(selected_source_ids) == 1:
            where_clause = {"file_id": selected_source_ids[0]}
        else:
            where_clause = {"file_id": {"$in": selected_source_ids}}
            
    documents = []
    metadatas = []
    sources_used = []
    context = ""
    
    try:
        if where_clause:
            results = collection.query(query_texts=[query], n_results=6, where=where_clause)
        else:
            results = collection.query(query_texts=[query], n_results=6)
            
        documents = results.get("documents", [[]])[0]
        metadatas = results.get("metadatas", [[]])[0]
        
        for doc, meta in zip(documents, metadatas):
            context += f"--- Source: {meta.get('source', 'Doc')} ---\n{doc}\n\n"
            if meta.get('source') and meta['source'] not in sources_used:
                sources_used.append(meta['source'])
    except Exception as e:
        print(f"Chroma DB query warning: {e}")

    system_prompt = (
        "You are StudySnap AI, an advanced educational AI pair programmer and study assistant.\n"
        "Use the provided source context retrieved from the user's uploaded materials to answer the user's question.\n"
        "CRITICAL INSTRUCTIONS FOR IMAGE & DIAGRAM QUESTIONS:\n"
        "1. When the user asks about an uploaded image or diagram (e.g. 'Explain me about the image I have provided'), base your explanation STRICTLY on the retrieved IMAGE OVERVIEW, VISUAL ELEMENTS, EXTRACTED TEXT, and DIAGRAM FLOW present in the context.\n"
        "2. Structure your explanation cleanly using Markdown headers (e.g., ### Overview, ### Main Components, ### How it works / Step-by-Step Flow, ### Extracted Text & Labels).\n"
        "3. Only include information that was actually identified from the uploaded image.\n"
        "4. Do NOT trigger web search or rely on generic unverified assumptions.\n\n"
        f"IMPORTANT: You must write your ENTIRE response in {response_language}.\n\n"
        f"Retrieved Source Context:\n{context}"
    )
    
    messages = [{"role": "system", "content": system_prompt}]
    for msg in history:
        messages.append({"role": msg.get("role", "user"), "content": msg.get("content", "")})
        
    final_query = query
    if response_language and response_language.lower() != "english":
        lang_def = ""
        if response_language.lower() == "hinglish":
            lang_def = " (a mix of Hindi and English written in the English alphabet)"
        elif response_language.lower() == "manglish":
            lang_def = " (a mix of Malayalam and English written in the English alphabet)"
            
        final_query += f"\n\n[CRITICAL DIRECTIVE: You MUST write your ENTIRE response in {response_language}{lang_def}.]"

    messages.append({"role": "user", "content": final_query})
    
    active_model = get_best_model_name(is_vision=False)
    try:
        response = await lm_studio_client.chat.completions.create(
            model=active_model,
            messages=messages,
            temperature=0.7,
            max_tokens=1200,
        )
        answer = response.choices[0].message.content
        if sources_used and "Source:" not in answer and "Sources:" not in answer:
            src_str = ", ".join(sources_used)
            answer += f"\n\n**Source:** {src_str}"
        return answer, sources_used
    except Exception as e:
        print(f"LM Studio LLM Exception/Fallback: {e}")
        if sympy_res and sympy_res.get("problem_type") != "Unknown / Unresolved":
            return format_sympy_result_to_markdown(sympy_res), ["SymPy Engine"]
        
        if context:
            return f"**Relevant Context Found from Documents & Image Analysis:**\n\n{context[:1500]}\n\n*(Note: LM Studio connection notice. Ensure LM Studio server is active on port 1234.)*", sources_used
            
        return "StudySnap AI active! Ask any question about your uploaded sources or images.", []

async def generate_podcast_script(selected_source_ids: list = None, response_language: str = "English"):
    where_clause = None
    if selected_source_ids and len(selected_source_ids) > 0:
        if len(selected_source_ids) == 1:
            where_clause = {"file_id": selected_source_ids[0]}
        else:
            where_clause = {"file_id": {"$in": selected_source_ids}}
            
    try:
        if where_clause:
            results = collection.get(where=where_clause, limit=20)
        else:
            results = collection.get(limit=20)
            
        documents = results.get("documents", [])
        metadatas = results.get("metadatas", [])
    except Exception as e:
        print(f"Chroma DB retrieval error in podcast script: {e}")
        documents, metadatas = [], []
        
    context = ""
    for doc, meta in zip(documents, metadatas):
        src_name = meta.get('source', 'Document') if meta else 'Document'
        context += f"--- Source: {src_name} ---\n{doc}\n\n"
        
    lang_def = ""
    if response_language and response_language.lower() == "hinglish":
        lang_def = " (a mix of Hindi and English written in the English alphabet)"
    elif response_language and response_language.lower() == "manglish":
        lang_def = " (a mix of Malayalam and English written in the English alphabet)"

    system_prompt = (
        "You are an expert podcast producer. Create a multi-turn, engaging podcast script between two hosts, "
        "Host A and Host B, discussing the provided documents. "
        f"IMPORTANT: The hosts MUST speak and banter in {response_language}{lang_def}. Write the actual dialogue in this language.\n"
        "Host A should introduce the topic, and they should have a natural back-and-forth banter. "
        "Return the output STRICTLY as a JSON array of objects, with each object having 'speaker' ('Host A' or 'Host B') "
        "and 'text' (the spoken dialogue). Do not include any other text outside the JSON array.\n\n"
        "Context:\n" + (context if context else "General study material.")
    )
    
    messages = [{"role": "system", "content": system_prompt}]
    active_model = get_best_model_name(is_vision=False)
    
    try:
        response = await lm_studio_client.chat.completions.create(
            model=active_model,
            messages=messages,
            temperature=0.7,
            max_tokens=2048,
        )
        raw_response = response.choices[0].message.content
        
        json_match = re.search(r'\[.*\]', raw_response, re.DOTALL)
        if json_match:
            script_json = json.loads(json_match.group(0))
        else:
            script_json = json.loads(raw_response)
        return script_json
    except Exception as e:
        print(f"LM Studio offline or JSON parse error in generate_podcast_script: {e}")
        ctx_snippet = context[:200].replace('\n', ' ') if context else 'the uploaded study materials'
        return [
            {"speaker": "Host A", "text": "Welcome to StudySnap Audio Overview!"}, 
            {"speaker": "Host B", "text": f"Today we are reviewing key concepts from our documents: {ctx_snippet}..."},
            {"speaker": "Host A", "text": "Make sure to explore the dashboard for complete details and study notes!"}
        ]

async def generate_tts_audio(script_json):
    audio_files = []
    voices = {
        "Host A": "en-US-ChristopherNeural",
        "Host B": "en-US-AriaNeural"
    }
    
    for i, line in enumerate(script_json):
        speaker = line.get("speaker", "Host A")
        text = line.get("text", "")
        if not text:
            continue
        voice = voices.get(speaker, voices["Host A"])
        
        output_file = f"temp_{uuid.uuid4()}_{i}.mp3"
        try:
            communicate = edge_tts.Communicate(text, voice)
            await communicate.save(output_file)
            audio_files.append(output_file)
        except Exception as tts_err:
            print(f"edge-tts generation error for line {i}: {tts_err}")
            
    if not audio_files:
        raise ValueError("Failed to generate any audio segments.")

    final_output_name = f"podcast_{uuid.uuid4()}.mp3"
    final_output = f"uploads/{final_output_name}"
    
    try:
        combined = AudioSegment.empty()
        for file in audio_files:
            try:
                segment = AudioSegment.from_mp3(file)
                combined += segment
                combined += AudioSegment.silent(duration=300)
            except Exception as e:
                print(f"Error merging segment {file}: {e}")
                
        if len(combined) > 0:
            combined.export(final_output, format="mp3")
        else:
            shutil.copy(audio_files[0], final_output)
    except Exception as pydub_err:
        print(f"Pydub audio export error: {pydub_err}. Using first segment as fallback.")
        shutil.copy(audio_files[0], final_output)
    
    for file in audio_files:
        if os.path.exists(file):
            try:
                os.remove(file)
            except Exception:
                pass
            
    return final_output_name

async def generate_infographic_image(selected_source_ids: list, style: str, detail_level: str, custom_prompt: str, response_language: str = "English"):
    where_clause = None
    if selected_source_ids and len(selected_source_ids) > 0:
        if len(selected_source_ids) == 1:
            where_clause = {"file_id": selected_source_ids[0]}
        else:
            where_clause = {"file_id": {"$in": selected_source_ids}}
            
    try:
        if where_clause:
            results = collection.get(where=where_clause, limit=20)
        else:
            results = collection.get(limit=20)
        documents = results.get("documents", [])
    except Exception:
        documents = []
        
    context = "\n\n".join(documents)
    
    lang_def = ""
    if response_language and response_language.lower() == "hinglish":
        lang_def = " (a mix of Hindi and English written in the English alphabet)"
    elif response_language and response_language.lower() == "manglish":
        lang_def = " (a mix of Malayalam and English written in the English alphabet)"

    system_prompt = (
        "You are an expert prompt engineer for an AI image generator (like Midjourney/DALL-E). "
        "Your task is to analyze the provided documents and create a highly detailed, descriptive visual prompt for an infographic. "
        f"The user wants a '{style}' style infographic with a '{detail_level}' level of detail. "
        f"User's custom instructions: {custom_prompt}\n"
        f"Any explicit text described in the image prompt should be in {response_language}{lang_def}.\n\n"
        "Extract the core concepts and data from the text, and describe exactly what the infographic should look like visually. "
        "Focus on layout, colors, icons, and visual metaphors. "
        "Return ONLY the image prompt text, no pleasantries or explanation."
    )
    
    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": f"Here is the source content:\n\n{context[:2000]}"}
    ]
    
    active_model = get_best_model_name(is_vision=False)
    image_prompt = ""
    try:
        response = await lm_studio_client.chat.completions.create(
            model=active_model,
            messages=messages,
            temperature=0.7,
            max_tokens=500,
        )
        image_prompt = response.choices[0].message.content.strip()
    except Exception as e:
        print(f"LM Studio offline for infographic prompt generation: {e}")
        image_prompt = f"Educational {style} infographic diagram showing {custom_prompt or 'key concepts and data visualization'}, vibrant colors, high quality 8k layout"

    encoded_prompt = urllib.parse.quote(image_prompt)
    image_url = f"https://image.pollinations.ai/prompt/{encoded_prompt}?width=1024&height=1024&nologo=true"
    
    final_output_name = f"infographic_{uuid.uuid4()}.png"
    final_output = f"uploads/{final_output_name}"
    
    try:
        async with aiohttp.ClientSession() as session:
            async with session.get(image_url, timeout=15) as resp:
                if resp.status == 200:
                    with open(final_output, 'wb') as f:
                        f.write(await resp.read())
                else:
                    raise ValueError(f"Pollinations API status: {resp.status}")
    except Exception as download_err:
        print(f"Image download error: {download_err}")
        from PIL import Image, ImageDraw
        img = Image.new('RGB', (1024, 1024), color='#1e1e2e')
        d = ImageDraw.Draw(img)
        d.text((100, 500), f"Infographic: {style}\nPrompt: {custom_prompt or 'Educational Overview'}", fill='#a3e5d8')
        img.save(final_output)

    return final_output_name

async def generate_mind_map_data(selected_source_ids: list, custom_prompt: str, response_language: str = "English"):
    where_clause = None
    if selected_source_ids and len(selected_source_ids) > 0:
        if len(selected_source_ids) == 1:
            where_clause = {"file_id": selected_source_ids[0]}
        else:
            where_clause = {"file_id": {"$in": selected_source_ids}}
            
    try:
        if where_clause:
            results = collection.get(where=where_clause, limit=20)
        else:
            results = collection.get(limit=20)
        documents = results.get("documents", [])
    except Exception:
        documents = []

    context = "\n\n".join(documents)
    
    lang_def = ""
    if response_language and response_language.lower() == "hinglish":
        lang_def = " (a mix of Hindi and English written in the English alphabet)"
    elif response_language and response_language.lower() == "manglish":
        lang_def = " (a mix of Malayalam and English written in the English alphabet)"

    system_prompt = (
        "You are an expert knowledge architect. Your task is to analyze the provided documents and create a structured Mind Map. "
        f"User's custom instructions: {custom_prompt}\n"
        f"IMPORTANT: The node labels and context descriptions MUST be in {response_language}{lang_def}.\n\n"
        "Return the output STRICTLY as a JSON object with two keys: 'nodes' and 'edges'. "
        "Each node should have: "
        "- id: string (unique identifier) "
        "- data: object containing 'label' (short title) and 'context' (1-2 sentences explaining this node based on sources) "
        "- position: object containing 'x' and 'y' coordinates (arrange them hierarchically, e.g., root at 250,0, children spread out below). "
        "Each edge should have: "
        "- id: string (e.g., 'e1-2') "
        "- source: string (id of parent node) "
        "- target: string (id of child node). "
        "Do not include any markdown formatting, only pure JSON."
    )
    
    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": f"Here is the source content:\n\n{context[:2000]}"}
    ]
    
    active_model = get_best_model_name(is_vision=False)
    try:
        response = await lm_studio_client.chat.completions.create(
            model=active_model,
            messages=messages,
            temperature=0.3,
            max_tokens=2000,
        )
        content = response.choices[0].message.content.strip()
        if content.startswith("```json"):
            content = content[7:-3].strip()
        elif content.startswith("```"):
            content = content[3:-3].strip()
        return json.loads(content)
    except Exception as e:
        print("LM Studio offline / Failed to parse JSON mind map:", e)
        nodes = [
            {"id": "1", "data": {"label": "StudySnap AI", "context": "Central Hub for your uploaded sources"}, "position": {"x": 250, "y": 50}}
        ]
        edges = []
        if documents:
            for idx, doc in enumerate(documents[:4]):
                nid = str(idx + 2)
                snippet = doc[:80].replace("\n", " ") + "..."
                nodes.append({
                    "id": nid,
                    "data": {"label": f"Source Topic {idx + 1}", "context": snippet},
                    "position": {"x": 100 + (idx * 200), "y": 200}
                })
                edges.append({"id": f"e1-{nid}", "source": "1", "target": nid})
        else:
            nodes.append({"id": "2", "data": {"label": "Study Overview", "context": "Upload files or links to generate dynamic mind maps"}, "position": {"x": 250, "y": 200}})
            edges.append({"id": "e1-2", "source": "1", "target": "2"})
            
        return {"nodes": nodes, "edges": edges}
